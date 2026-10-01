"""人間確認用の売上 Excel（売上・請求書集計（暫定））のテスト。すべて一時フォルダ上の完全な架空データ。"""

from __future__ import annotations

import json
from collections import Counter

import pytest
from openpyxl import load_workbook

from blue_return_ai import sales_workbook
from blue_return_ai.annual import run
from blue_return_ai.sample_template_c import write_sample as write_c
from blue_return_ai.sample_template_instructor import write_sample as write_instructor

CONFIG = {
    "version": 1,
    "sources": {"invoices": "data/invoices"},
    "processors": {"invoices": {"templates": ["template_c", "template_instructor"]}},
}
WORKBOOK = sales_workbook.workbook_file_name(2026)


@pytest.fixture
def result(tmp_path):
    repo = tmp_path / "repo"
    invoices = repo / "data" / "invoices"
    write_c(invoices / "a.xlsx", pattern="A")                       # 税込 77000、明細 2 件
    write_c(invoices / "e.xlsx", pattern="E")                       # 税込 11000、明細 1 件
    write_c(invoices / "d.xlsx", pattern="D")                       # 税込は不明（null）
    write_c(invoices / "x.xlsx", pattern="A", overrides={"B3": "架空第二商事株式会社"})
    write_instructor(invoices / "i.xlsx")                           # 税込は不明（null）、明細 2 件
    config = repo / "data" / "config" / "annual.json"
    config.parent.mkdir(parents=True)
    config.write_text(json.dumps(CONFIG), encoding="utf-8")
    return run(2026, config, repo / "data" / "output" / "annual", repo_root=repo)


@pytest.fixture
def workbook(result):
    return load_workbook(result["run_dir"] / WORKBOOK)


def summary_values(workbook) -> dict[str, tuple]:
    sheet = workbook["集計"]
    return {row[0]: (row[1], row[2]) for row in sheet.iter_rows(min_row=6, values_only=True) if row[0]}


def table(workbook, name) -> tuple[list, list[dict]]:
    sheet = workbook[name]
    header = [c.value for c in sheet[3]]
    rows = [dict(zip(header, values)) for values in sheet.iter_rows(min_row=4, values_only=True)]
    return header, rows


def test_workbook_is_created_in_run_dir_with_four_sheets(result, workbook):
    assert (result["run_dir"] / WORKBOOK).exists()
    assert WORKBOOK in result["manifest"]["outputs"]
    assert workbook.sheetnames == ["集計", "売上一覧", "売上明細", "要確認"]
    for name in workbook.sheetnames:
        assert "2026年 売上・請求書集計（暫定）" in workbook[name]["A1"].value
        assert "暫定資料" in workbook[name]["A2"].value


def test_sales_list(workbook):
    header, rows = table(workbook, "売上一覧")
    assert len(rows) == 5
    for column in ("資料日付（請求日等）", "売上計上日", "業務期間（開始）", "取引先", "内容", "税抜金額", "消費税額",
                   "税込金額", "源泉徴収税額", "その他控除", "振込予定額", "支払期限", "税区分", "源泉徴収の状態",
                   "確認状態（review_status）", "警告", "テンプレート", "帳票パターン", "record_id"):
        assert column in header
    # 売上計上日は請求日で補完しない
    assert all(r["売上計上日"] is None for r in rows)
    assert all(r["売上計上日の状態"] == "unconfirmed" for r in rows)
    pattern_a = next(r for r in rows if r["帳票パターン"] == "A")
    assert pattern_a["税込金額"] == 77000
    assert pattern_a["資料日付（請求日等）"].date().isoformat() == "2026-05-31"
    # 不明な金額は空欄（0 にしない）
    pattern_d = next(r for r in rows if r["帳票パターン"] == "D")
    assert pattern_d["税込金額"] is None
    assert pattern_d["小計（意味未確定・確認用）"] == 8000
    instructor = next(r for r in rows if r["帳票パターン"] == "instructor")
    assert instructor["税込金額"] is None
    assert instructor["振込予定額"] is None
    assert instructor["ご請求金額（帳票の値・確認用）"] == 12973


def test_text_starting_with_equal_is_not_a_formula(tmp_path):
    sales = [{"record_id": "r1", "source_key": "k1", "customer": "=1+1", "description": "+架空",
              "review_status": "needs_review"}]
    issues = [{"source_key": "k1", "message": "=HYPERLINK(\"x\")", "issue_code": "FIELD_MISSING"}]
    path = sales_workbook.write_workbook(tmp_path / WORKBOOK, sales_workbook.build_workbook(
        sales, [], issues, Counter(), 2026, "run", "now"))
    workbook = load_workbook(path)
    sheet = workbook["売上一覧"]
    header = [c.value for c in sheet[3]]
    customer = sheet.cell(row=4, column=header.index("取引先") + 1)
    assert customer.value == "=1+1"
    assert customer.data_type == "s"
    issue_sheet = workbook["要確認"]
    message = issue_sheet.cell(row=4, column=[c.value for c in issue_sheet[3]].index("内容") + 1)
    assert message.data_type == "s"


def test_line_items_are_linked_by_sales_no(workbook):
    _, sales = table(workbook, "売上一覧")
    _, lines = table(workbook, "売上明細")
    assert len(lines) == 2 + 1 + 1 + 2 + 2
    by_no = {s["売上No"]: s for s in sales}
    for line in lines:
        assert by_no[line["売上No"]]["record_id"] == line["record_id"]
    instructor_lines = [l for l in lines if by_no[l["売上No"]]["帳票パターン"] == "instructor"]
    assert [l["税率(%)"] for l in instructor_lines] == [10, 10]
    assert [l["単位"] for l in instructor_lines] == ["時間", "回"]


def test_summary_does_not_treat_null_as_zero(workbook):
    values = summary_values(workbook)
    assert values["取り込み済み売上資料（件）"][0] == 5
    assert values["売上計上日 確定（件）"][0] == 0
    assert values["売上計上日 未確定（件）"][0] == 5
    # 税込金額：A 2 件（77000×2）＋ E（11000）が判明、D と講師形式は不明
    total, note = values["税込金額：判明分の合計（円）"]
    assert total == 77000 * 2 + 11000
    assert "5 件中 3 件の合計" in note
    assert "2 件は値が不明" in note
    assert values["税込金額：値が不明（件）"][0] == 2
    assert values["review_status = needs_review（件）"][0] == 5
    assert values["review_status = reviewed（件）"][0] == 0
    assert values["帳票パターン = instructor（件）"][0] == 1


def test_issue_sheet_contains_only_sales_issues(result, workbook):
    _, issues = table(workbook, "要確認")
    expected = result["manifest"]["processing"]["unresolved"]["total"]
    assert len(issues) == expected  # この構成では要確認はすべて売上のもの
    assert sum(i["警告コード"] == "REVENUE_DATE_UNCONFIRMED" for i in issues) == 5
    assert all(i["売上No"] for i in issues)
    assert all(i["判断（記入欄）"] is None for i in issues)


def test_summary_rows_with_no_unknown_values():
    sales = [{"gross_amount": 1000, "review_status": "unreviewed", "revenue_date": None},
             {"gross_amount": 2000, "review_status": "unreviewed", "revenue_date": None}]
    rows = {r[0]: (r[1], r[2]) for r in sales_workbook.build_summary_rows(sales, Counter())}
    assert rows["税込金額：判明分の合計（円）"] == (3000, "2 件すべての値が判明")
    assert rows["税抜金額：判明分の合計（円）"][0] is None  # 判明分が無ければ合計も空欄（0 にしない）
    assert rows["税抜金額：値が不明（件）"][0] == 2


def test_empty_sales_produce_valid_workbook(tmp_path):
    workbook = sales_workbook.build_workbook([], [], [], Counter(), 2026, "run", "now")
    path = sales_workbook.write_workbook(tmp_path / WORKBOOK, workbook)
    assert load_workbook(path).sheetnames == ["集計", "売上一覧", "売上明細", "要確認"]


def test_existing_workbook_is_not_overwritten(tmp_path):
    path = tmp_path / WORKBOOK
    path.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        sales_workbook.write_workbook(path, sales_workbook.build_workbook(
            [], [], [], Counter(), 2026, "run", "now"))
    assert path.read_bytes() == b"existing"
