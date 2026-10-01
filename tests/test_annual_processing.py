"""年間処理（売上・カード・銀行の取込の統合）のテスト。一時フォルダ上の完全な架空データのみ。"""

from __future__ import annotations

import csv
import json

import pytest
from openpyxl import Workbook

from blue_return_ai import annual
from blue_return_ai.annual import run
from blue_return_ai.config import ConfigError
from blue_return_ai.sample_template_c import write_sample
from test_transactions_csv import (
    BANK_HEADER,
    BANK_PREAMBLE,
    BANK_ROWS,
    CARD_HEADER,
    CARD_PREAMBLE,
    CARD_ROWS,
    write_csv,
)

CONFIG = {
    "version": 1,
    "sources": {
        "invoices": "data/invoices",
        "bank": "data/bank",
        "credit_card": "data/credit-card",
        "household": "data/household",
    },
    "decisions_dir": "data/decisions",
    "processors": {
        "invoices": {"template": "template_c"},
        "credit_card": {"format": "card_csv_a"},
        "bank": {"format": "bank_csv_a"},
    },
}

FILE_NAMES = ["架空請求_取引先甲.xlsx", "架空請求_取引先乙.xlsx", "架空講師.xlsx", "架空支払明細.pdf",
              "架空カード_2026年7月.csv", "架空銀行_6月.csv", "架空銀行_6月後半.csv", "架空銀行_壊れ.csv"]


def write_config(repo, data=CONFIG):
    path = repo / "data" / "config" / "annual.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return path


@pytest.fixture
def repo(tmp_path):
    repo = tmp_path / "repo"
    invoices = repo / "data" / "invoices"
    write_sample(invoices / FILE_NAMES[0])
    write_sample(invoices / FILE_NAMES[1], items=[("架空単発業務", 1, 50000, 50000)],
                 summary={"net_amount": 50000, "tax_amount": 5000, "gross_amount": 55000,
                          "withholding_tax": 5105, "expected_payment_amount": 49895})
    other = Workbook()
    other.active.title = "架空の別書式"
    other.save(invoices / FILE_NAMES[2])
    (invoices / FILE_NAMES[3]).write_bytes(b"%PDF-fake")

    card_dir = repo / "data" / "credit-card"
    card_dir.mkdir(parents=True)
    write_csv(card_dir / FILE_NAMES[4], CARD_PREAMBLE, CARD_HEADER, CARD_ROWS)

    bank_dir = repo / "data" / "bank"
    bank_dir.mkdir(parents=True)
    write_csv(bank_dir / FILE_NAMES[5], BANK_PREAMBLE, BANK_HEADER, BANK_ROWS)
    # 期間の重なる CSV（2〜4 行目と同じ取引。明細通番は振り直されている）
    overlap = ["10" + row[1:] for row in BANK_ROWS[1:4]]
    write_csv(bank_dir / FILE_NAMES[6], BANK_PREAMBLE, BANK_HEADER, overlap)
    write_csv(bank_dir / FILE_NAMES[7], BANK_PREAMBLE, "日付,金額", BANK_ROWS)
    write_config(repo)
    return repo


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def run_repo(repo):
    return run(2026, repo / "data" / "config" / "annual.json",
               repo / "data" / "output" / "annual", repo_root=repo)


def test_all_outputs_are_created(repo):
    run_dir = run_repo(repo)["run_dir"]
    assert sorted(p.name for p in run_dir.iterdir()) == sorted(annual.output_files(2026))
    assert run_dir.parent == (repo / "data" / "output" / "annual").resolve()


def test_sales_are_imported_from_template_c_only(repo):
    run_dir = run_repo(repo)["run_dir"]
    sales = read_csv(run_dir / "sales.csv")
    assert len(sales) == 2
    assert {s["source_name"] for s in sales} == {"template_c"}
    assert {s["gross_amount"] for s in sales} == {"77000", "55000"}
    assert all(s["source_key"].startswith("sales-1-") for s in sales)
    assert all(s["revenue_date"] == "" and s["revenue_date_status"] == "unconfirmed" for s in sales)
    assert len(read_csv(run_dir / "sales_line_items.csv")) == 3

    inventory = {r["relative_path"].split("/")[-1]: r for r in read_csv(run_dir / "documents_inventory.csv")}
    assert inventory[FILE_NAMES[0]]["status"] == "imported"
    assert inventory[FILE_NAMES[2]]["status"] == "unresolved"
    assert inventory[FILE_NAMES[2]]["note"] == "template_sheet_not_found"
    assert inventory[FILE_NAMES[3]]["status"] == "unresolved"  # PDF からは売上を作らない


def test_card_and_bank_transactions(repo):
    run_dir = run_repo(repo)["run_dir"]
    card = read_csv(run_dir / "card_transactions.csv")
    assert len(card) == 6
    assert all(c["business_classification"] == "unknown" for c in card)
    assert all(c["review_status"] == "needs_review" for c in card)
    # 同じ内容の 2 件は消さずに重複候補として残す
    assert sum("DUPLICATE_CANDIDATE" in c["warnings"] for c in card) == 2

    bank = read_csv(run_dir / "bank_transactions.csv")
    active = [b for b in bank if b["record_status"] == "active"]
    duplicates = [b for b in bank if b["record_status"] == "duplicate"]
    assert len(active) == 6
    assert len(duplicates) == 3  # 期間の重なる CSV の 3 件は二重登録しない
    assert all(d["duplicate_of_document"] for d in duplicates)
    zero = next(b for b in active if b["description"] == "架空手数料")
    assert zero["withdrawal_amount"] == "0"
    assert zero["deposit_amount"] == ""

    inventory = {r["relative_path"].split("/")[-1]: r for r in read_csv(run_dir / "documents_inventory.csv")}
    assert inventory[FILE_NAMES[4]]["status"] == "imported"
    assert inventory[FILE_NAMES[6]]["status"] == "imported"
    assert inventory[FILE_NAMES[7]]["status"] == "unresolved"
    assert inventory[FILE_NAMES[7]]["note"] == "format_mismatch"


def test_unresolved_items_are_integrated(repo):
    run_dir = run_repo(repo)["run_dir"]
    items = read_csv(run_dir / "unresolved_items.csv")
    assert list(items[0]) == annual.UNRESOLVED_COLUMNS
    by_type = {}
    for item in items:
        by_type.setdefault(item["source_type"], set()).add(item["issue_code"])
    assert {"REVENUE_DATE_UNCONFIRMED", "FIELD_MISSING"} <= by_type["sales"]
    assert {"CLASSIFICATION_UNKNOWN", "DUPLICATE_CANDIDATE", "CARD_NEGATIVE_AMOUNT",
            "ROW_NOT_TRANSACTION"} <= by_type["card"]
    assert {"CLASSIFICATION_UNKNOWN", "CSV_HEADER_MISMATCH", "BANK_AMOUNT_MISSING"} <= by_type["bank"]
    assert all(item["message"] and item["decision_status"] == "" for item in items)
    # 重複として除外した取引は要確認に出さない（有効な銀行取引 6 件分の分類未判断のみ）
    assert sum(i["issue_code"] == "CLASSIFICATION_UNKNOWN" and i["source_type"] == "bank"
               for i in items) == 6
    # 要確認一覧に実データの値・ファイル名を書かない
    text = (run_dir / "unresolved_items.csv").read_text(encoding="utf-8-sig")
    for value in ("架空", "1200", "50000", "77000") + tuple(FILE_NAMES):
        assert value not in text


def test_manifest_counts(repo):
    manifest = run_repo(repo)["manifest"]
    processing = manifest["processing"]
    assert manifest["processors"] == {"invoices": "template_c", "bank": "bank_csv_a",
                                      "credit_card": "card_csv_a", "household": "not_configured"}
    assert processing["files"]["invoices"] == {
        "files": 2 + 1, "imported": 2, "failed": 0, "template_not_matched": 1,
        "duplicate_files": 0, "skipped_by_decision": 0}
    assert processing["files"]["bank"]["failed"] == 1
    assert processing["sales"] == {"records": 2, "needs_review": 2,
                                   "by_template_pattern": {"A": 2}}
    assert processing["card"]["transactions"] == 6
    assert processing["card"]["rows_not_transaction"] == 2
    assert processing["bank"]["transactions"] == 6
    assert processing["bank"]["duplicate_transactions"] == 3
    assert processing["unresolved"]["total"] == sum(processing["unresolved"]["by_code"].values())


def test_rerun_regenerates_same_results_in_new_directory(repo):
    first = run_repo(repo)
    second = run_repo(repo)
    assert first["run_dir"] != second["run_dir"]
    for name in ("sales.csv", "card_transactions.csv", "bank_transactions.csv"):
        rows_1 = read_csv(first["run_dir"] / name)
        rows_2 = read_csv(second["run_dir"] / name)
        assert [r["source_key"] for r in rows_1] == [r["source_key"] for r in rows_2]
    assert first["manifest"]["processing"] == second["manifest"]["processing"]


def test_same_file_twice_is_not_imported_twice(repo):
    invoices = repo / "data" / "invoices"
    (invoices / "copy").mkdir()
    (invoices / "copy" / "コピー.xlsx").write_bytes((invoices / FILE_NAMES[0]).read_bytes())
    result = run_repo(repo)
    assert len(read_csv(result["run_dir"] / "sales.csv")) == 2
    assert result["manifest"]["processing"]["files"]["invoices"]["duplicate_files"] == 1


def test_human_decision_skips_processing(repo):
    import hashlib
    sha = hashlib.sha256((repo / "data" / "invoices" / FILE_NAMES[1]).read_bytes()).hexdigest()
    decisions = repo / "data" / "decisions" / "documents.csv"
    decisions.parent.mkdir(parents=True)
    decisions.write_text(f"sha256,status,note\n{sha},manual,\n", encoding="utf-8")
    result = run_repo(repo)
    assert len(read_csv(result["run_dir"] / "sales.csv")) == 1
    assert result["manifest"]["processing"]["files"]["invoices"]["skipped_by_decision"] == 1


@pytest.mark.parametrize(
    "processors",
    [
        {"invoices": {"template": "template_x"}},
        {"bank": {"format": "bank_csv_z"}},
        {"bank": {"format": "bank_csv_a", "options": {"delimiter": "tab"}}},
        {"bank": {"format": "bank_csv_a", "options": {"header_row": "13"}}},
        {"credit_card": {"template": "card_csv_a"}},
        {"invoices": {"template": "template_c", "options": {"x": 1}}},
        {"household": {"format": "x"}},
    ],
)
def test_invalid_processors_are_rejected(repo, processors):
    write_config(repo, {**CONFIG, "processors": processors})
    with pytest.raises(ConfigError):
        run_repo(repo)
    assert not (repo / "data" / "output").exists()


def test_cli_prints_counts_only(repo, monkeypatch, capsys):
    monkeypatch.setattr(annual, "REPO_ROOT", repo)
    code = annual.main(["--year", "2026", "--config", str(repo / "data" / "config" / "annual.json"),
                        "--out-root", str(repo / "data" / "output" / "annual")])
    out = capsys.readouterr()
    assert code == 0
    text = out.out + out.err
    assert "売上: 2 件（要確認 2 件）" in text
    assert "カード取引: 6 件" in text
    assert "銀行取引: 6 件" in text
    assert "CLASSIFICATION_UNKNOWN:" in text
    # 表示してよいのは、プログラムが決めた出力ファイル名（人間確認用 Excel）だけ
    assert "sales_summary_2026_provisional.xlsx" in text
    text = text.replace("sales_summary_2026_provisional.xlsx", "")
    for value in ("架空", "1200", "50000", "77000", ".xlsx", ".csv") + tuple(FILE_NAMES):
        assert value not in text


def test_mixed_invoice_patterns_and_serial_dates(tmp_path):
    from datetime import datetime

    from openpyxl.utils.datetime import to_excel
    repo = tmp_path / "repo"
    invoices = repo / "data" / "invoices"
    serial = to_excel(datetime(2026, 4, 30))
    for pattern in "ABCD":
        write_sample(invoices / f"架空_{pattern}.xlsx", pattern=pattern, document_date=serial)
    write_config(repo, {**CONFIG, "processors": {"invoices": {"template": "template_c"}}})
    result = run_repo(repo)
    sales = read_csv(result["run_dir"] / "sales.csv")
    assert sorted(s["template_pattern"] for s in sales) == ["A", "B", "C", "D"]
    assert {s["document_date"] for s in sales} == {"2026-04-30"}
    by_pattern = {s["template_pattern"]: s for s in sales}
    assert by_pattern["B"]["withholding_base"] == "30000"
    assert by_pattern["D"]["document_subtotal"] == "8000"
    assert by_pattern["D"]["gross_amount"] == ""
    assert result["manifest"]["processing"]["sales"]["by_template_pattern"] == {
        "A": 1, "B": 1, "C": 1, "D": 1}
    codes = result["manifest"]["processing"]["unresolved"]["by_code"]
    assert "DATE_INVALID" not in codes
    assert "TEMPLATE_LABEL_MISMATCH" not in codes


# ---------------------------------------------------------------- 複数テンプレートの振り分け

MULTI_TEMPLATES = {**CONFIG, "processors": {
    "invoices": {"templates": ["template_c", "template_instructor"]}}}


def test_invoices_are_dispatched_by_sheet_name(tmp_path):
    from openpyxl import Workbook

    from blue_return_ai.sample_template_instructor import write_sample as write_instructor
    repo = tmp_path / "repo"
    invoices = repo / "data" / "invoices"
    for pattern in "ABCDE":
        write_sample(invoices / f"架空_{pattern}.xlsx", pattern=pattern)
    write_instructor(invoices / "架空_講師.xlsx")
    # 両方のシートを持つブック（振り分けを決められない）
    both = Workbook()
    both.active.title = "misoca_invoice"
    both.create_sheet("講師")
    both.save(invoices / "架空_両方.xlsx")
    # どちらのシートも無いブック
    other = Workbook()
    other.active.title = "架空の別書式"
    other.save(invoices / "架空_別.xlsx")
    write_config(repo, MULTI_TEMPLATES)

    result = run_repo(repo)
    processing = result["manifest"]["processing"]
    assert result["manifest"]["processors"]["invoices"] == "template_c,template_instructor"
    assert processing["files"]["invoices"]["imported"] == 6
    assert processing["files"]["invoices"]["template_not_matched"] == 2
    assert processing["sales"]["by_template_pattern"] == {
        "A": 1, "B": 1, "C": 1, "D": 1, "E": 1, "instructor": 1}
    assert "TEMPLATE_LABEL_MISMATCH" not in processing["unresolved"]["by_code"]
    assert "DATE_INVALID" not in processing["unresolved"]["by_code"]

    inventory = {r["relative_path"].split("/")[-1]: r
                 for r in read_csv(result["run_dir"] / "documents_inventory.csv")}
    assert inventory["架空_講師.xlsx"]["status"] == "imported"
    assert inventory["架空_両方.xlsx"]["note"] == "template_ambiguous"
    assert inventory["架空_両方.xlsx"]["status"] == "unresolved"
    assert inventory["架空_別.xlsx"]["note"] == "template_sheet_not_found"

    sales = {s["template_pattern"]: s for s in read_csv(result["run_dir"] / "sales.csv")}
    assert sales["instructor"]["source_name"] == "template_instructor"
    assert sales["instructor"]["document_total"] == "12973"
    assert sales["instructor"]["expected_payment_amount"] == ""


def test_instructor_sheet_with_wrong_structure_is_counted_separately(tmp_path):
    from blue_return_ai.sample_template_instructor import write_sample as write_instructor
    repo = tmp_path / "repo"
    write_instructor(repo / "data" / "invoices" / "架空_講師.xlsx", overrides={"I20": "数量"})
    write_config(repo, MULTI_TEMPLATES)
    processing = run_repo(repo)["manifest"]["processing"]
    assert processing["sales"]["by_template_pattern"] == {"instructor_unidentified": 1}
    assert processing["unresolved"]["by_code"]["TEMPLATE_LABEL_MISMATCH"] == 2


@pytest.mark.parametrize(
    "invoices",
    [
        {"templates": []},
        {"templates": ["template_c", "template_c"]},
        {"templates": ["template_c", "template_unknown"]},
        {"templates": "template_c"},
        {"template": "template_c", "templates": ["template_instructor"]},
    ],
)
def test_invalid_templates_setting_is_rejected(repo, invoices):
    write_config(repo, {**CONFIG, "processors": {"invoices": invoices}})
    with pytest.raises(ConfigError):
        run_repo(repo)
