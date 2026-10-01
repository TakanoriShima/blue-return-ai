"""会計用の売上金額（税込経理）accounting_sales_amount のテスト。すべて完全な架空データ。"""

from __future__ import annotations

import csv
import json

import pytest
from openpyxl import load_workbook

from blue_return_ai import sales_workbook
from blue_return_ai.annual import run
from blue_return_ai.annual_processors import decide_accounting_sales_amount
from blue_return_ai.config import ConfigError, load_config
from blue_return_ai.sample_template_c import write_sample as write_c
from blue_return_ai.sample_template_instructor import write_sample as write_instructor
from conftest import REPO_ROOT

EXEMPT_INCLUSIVE = {"tax_status": "exempt", "invoice_registration": False,
                    "consumption_tax_accounting": "inclusive"}
CONFIG = {
    "version": 1,
    "sources": {"invoices": "data/invoices"},
    "processors": {"invoices": {"templates": ["template_c", "template_instructor"]}},
    "accounting": {"2026": EXEMPT_INCLUSIVE},
}
WORKBOOK = sales_workbook.workbook_file_name(2026)


# ---------------------------------------------------------------- 算出ルール（単体）

def record(**values):
    base = {"source_name": "template_c", "gross_amount": None, "net_amount": None, "tax_amount": None}
    return {**base, **values}


def test_gross_amount_is_used_when_known():
    assert decide_accounting_sales_amount(record(gross_amount=11000), {"template_pattern": "A"},
                                          EXEMPT_INCLUSIVE) == (11000, "gross_amount")


def test_known_zero_is_kept_as_zero():
    assert decide_accounting_sales_amount(record(gross_amount=0), {}, EXEMPT_INCLUSIVE) == (0, "gross_amount")
    assert decide_accounting_sales_amount(record(), {"template_pattern": "D", "document_total": 0},
                                          EXEMPT_INCLUSIVE) == (0, "document_total")


def test_pattern_d_uses_document_total():
    assert decide_accounting_sales_amount(record(), {"template_pattern": "D", "document_total": 8000},
                                          EXEMPT_INCLUSIVE) == (8000, "document_total")


@pytest.mark.parametrize(
    ("rec", "reference"),
    [
        (record(), {"template_pattern": "D", "document_total": None}),       # ご請求金額も不明
        (record(), {"template_pattern": "B", "document_total": 33000}),      # D 以外の「ご請求金額」は使わない
        (record(source_name="template_instructor"),                          # 講師形式の「合計」は意味未確定
         {"template_pattern": "instructor", "document_total": 12973}),
        (record(source_name="template_a"), {"template_pattern": None, "document_total": 5000}),
        (record(), {"template_pattern": None}),                              # パターン判定不可
    ],
)
def test_undecidable_amount_is_null(rec, reference):
    assert decide_accounting_sales_amount(rec, reference, EXEMPT_INCLUSIVE) == (None, None)


def instructor(**values):
    defaults = {"source_name": "template_instructor", "net_amount": 13000, "tax_amount": 1300,
                "warnings": []}
    return record(**{**defaults, **values})


INSTRUCTOR_REF = {"template_pattern": "instructor", "document_total": 12973}


def test_instructor_subtotal_is_used():
    rec = instructor()
    before = dict(rec)
    assert decide_accounting_sales_amount(rec, INSTRUCTOR_REF, EXEMPT_INCLUSIVE) == (13000, "instructor_subtotal")
    assert rec == before  # 原資料の値は変更しない


def test_instructor_final_total_is_not_used():
    amount, _ = decide_accounting_sales_amount(instructor(), INSTRUCTOR_REF, EXEMPT_INCLUSIVE)
    assert amount != INSTRUCTOR_REF["document_total"]


def test_instructor_zero_subtotal_is_kept():
    assert decide_accounting_sales_amount(instructor(net_amount=0), INSTRUCTOR_REF,
                                          EXEMPT_INCLUSIVE) == (0, "instructor_subtotal")


@pytest.mark.parametrize(
    ("rec", "reference"),
    [
        (instructor(), {"template_pattern": "instructor_unidentified"}),     # 構造を確認できない
        (instructor(net_amount=None), INSTRUCTOR_REF),                      # 小計が読めない
        (instructor(warnings=[{"code": "FORMULA_NO_CACHED_VALUE", "field": "net_amount"}]), INSTRUCTOR_REF),
        (record(source_name="template_c", net_amount=13000), INSTRUCTOR_REF),  # 講師形式以外の net_amount
        (record(source_name="template_c", net_amount=13000), {"template_pattern": "A"}),
    ],
)
def test_instructor_subtotal_requires_confirmed_structure(rec, reference):
    assert decide_accounting_sales_amount(rec, reference, EXEMPT_INCLUSIVE) == (None, None)


def test_instructor_rule_needs_exempt_inclusive():
    taxable = {**EXEMPT_INCLUSIVE, "tax_status": "taxable"}
    assert decide_accounting_sales_amount(instructor(), INSTRUCTOR_REF, taxable) == (None, None)


@pytest.mark.parametrize(
    "settings",
    [None, {},
     {**EXEMPT_INCLUSIVE, "tax_status": "taxable"},
     {**EXEMPT_INCLUSIVE, "consumption_tax_accounting": "exclusive"}],
)
def test_only_exempt_inclusive_is_supported(settings):
    assert decide_accounting_sales_amount(record(gross_amount=11000), {}, settings) == (None, None)


# ---------------------------------------------------------------- 設定ファイル

def write_config(repo, data):
    path = repo / "data" / "config" / "annual.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def without_accounting():
    return {k: v for k, v in CONFIG.items() if k != "accounting"}


def test_accounting_settings_are_loaded_by_year(tmp_path):
    repo = tmp_path / "repo"
    config = load_config(write_config(repo, {**CONFIG, "accounting": {
        "2026": EXEMPT_INCLUSIVE,
        "2027": {"tax_status": "taxable", "invoice_registration": True,
                 "consumption_tax_accounting": "exclusive"},
    }}), repo)
    assert config.accounting[2026] == EXEMPT_INCLUSIVE
    assert config.accounting[2027]["tax_status"] == "taxable"
    assert load_config(write_config(repo, without_accounting()), repo).accounting == {}


@pytest.mark.parametrize(
    "accounting",
    [
        [],
        {"26": EXEMPT_INCLUSIVE},
        {"2026": {**EXEMPT_INCLUSIVE, "tax_status": "free"}},
        {"2026": {**EXEMPT_INCLUSIVE, "invoice_registration": "no"}},
        {"2026": {**EXEMPT_INCLUSIVE, "consumption_tax_accounting": "net"}},
        {"2026": {"tax_status": "exempt"}},
        {"2026": {**EXEMPT_INCLUSIVE, "extra": 1}},
    ],
)
def test_invalid_accounting_settings_are_rejected(tmp_path, accounting):
    repo = tmp_path / "repo"
    with pytest.raises(ConfigError):
        load_config(write_config(repo, {**CONFIG, "accounting": accounting}), repo)


def test_example_config_has_2026_settings():
    config = load_config(REPO_ROOT / "sample_data" / "config" / "annual.example.json", REPO_ROOT)
    assert config.accounting[2026] == EXEMPT_INCLUSIVE


# ---------------------------------------------------------------- annual・CSV・Excel

@pytest.fixture
def repo(tmp_path):
    repo = tmp_path / "repo"
    invoices = repo / "data" / "invoices"
    write_c(invoices / "a.xlsx", pattern="A")                      # 税込 77000
    write_c(invoices / "z.xlsx", pattern="C", items=[("架空無償対応", 1, 0, 0)],
            summary={"net_amount": 0, "tax_amount": 0, "gross_amount": 0})  # 税込 0（既知の 0）
    write_c(invoices / "d.xlsx", pattern="D")                      # ご請求金額 8000
    write_instructor(invoices / "i.xlsx")                          # 講師形式：小計 13000（源泉徴収前）
    write_instructor(invoices / "u.xlsx", overrides={"I20": "数量"})  # 講師形式だが構造を確認できない → 決められない
    write_config(repo, CONFIG)
    return repo


def run_repo(repo):
    return run(2026, repo / "data" / "config" / "annual.json", repo / "data" / "output" / "annual",
               repo_root=repo)


def read_sales(run_dir):
    with (run_dir / "sales.csv").open(encoding="utf-8-sig", newline="") as f:
        return {row["template_pattern"]: row for row in csv.DictReader(f)}


def test_sales_csv_has_accounting_amount_and_keeps_source_values(repo):
    result = run_repo(repo)
    sales = read_sales(result["run_dir"])
    assert sales["A"]["accounting_sales_amount"] == "77000"
    assert sales["A"]["accounting_sales_amount_basis"] == "gross_amount"
    assert sales["A"]["gross_amount"] == "77000"                  # 原資料の値はそのまま
    assert sales["C"]["accounting_sales_amount"] == "0"           # 既知の 0 は 0
    assert sales["D"]["accounting_sales_amount"] == "8000"
    assert sales["D"]["accounting_sales_amount_basis"] == "document_total"
    # パターン D の原資料の金額は作らない・変えない
    assert sales["D"]["gross_amount"] == ""
    assert sales["D"]["net_amount"] == ""
    assert sales["D"]["tax_amount"] == ""
    assert sales["D"]["document_subtotal"] == "8000"
    # 講師形式：帳票の「小計」（源泉徴収前）を使い、最終の「合計」（源泉徴収後 12973）は使わない
    assert sales["instructor"]["accounting_sales_amount"] == "13000"
    assert sales["instructor"]["accounting_sales_amount_basis"] == "instructor_subtotal"
    assert sales["instructor"]["document_total"] == "12973"
    # 講師形式の原資料の値は変えない・作らない
    assert sales["instructor"]["net_amount"] == "13000"
    assert sales["instructor"]["tax_amount"] == "1300"
    assert sales["instructor"]["gross_amount"] == ""
    # 構造を確認できない講師形式は決めない
    assert sales["instructor_unidentified"]["accounting_sales_amount"] == ""
    assert sales["instructor_unidentified"]["accounting_sales_amount_basis"] == ""
    assert result["manifest"]["accounting"] == EXEMPT_INCLUSIVE


def test_undecidable_amount_is_reviewed(repo):
    result = run_repo(repo)
    by_code = result["manifest"]["processing"]["unresolved"]["by_code"]
    assert by_code["ACCOUNTING_SALES_AMOUNT_UNKNOWN"] == 1
    sales = read_sales(result["run_dir"])
    unidentified = sales["instructor_unidentified"]
    assert "ACCOUNTING_SALES_AMOUNT_UNKNOWN:accounting_sales_amount" in unidentified["warnings"]
    assert unidentified["review_status"] == "needs_review"
    for pattern in ("A", "C", "D", "instructor"):
        assert "ACCOUNTING_SALES_AMOUNT_UNKNOWN" not in sales[pattern]["warnings"]


def test_without_settings_all_amounts_are_unknown(repo):
    write_config(repo, without_accounting())
    result = run_repo(repo)
    assert all(row["accounting_sales_amount"] == "" for row in read_sales(result["run_dir"]).values())
    assert result["manifest"]["processing"]["unresolved"]["by_code"]["ACCOUNTING_SALES_AMOUNT_UNKNOWN"] == 5
    assert result["manifest"]["accounting"] is None


def test_excel_shows_accounting_amount_and_counts(repo):
    result = run_repo(repo)
    workbook = load_workbook(result["run_dir"] / WORKBOOK)
    sheet = workbook["売上一覧"]
    header = [c.value for c in sheet[3]]
    column = header.index("会計用売上金額（税込経理）") + 1
    pattern = header.index("帳票パターン") + 1
    values = {sheet.cell(row=r, column=pattern).value: sheet.cell(row=r, column=column).value
              for r in range(4, sheet.max_row + 1)}
    assert values == {"A": 77000, "C": 0, "D": 8000, "instructor": 13000, "instructor_unidentified": None}

    summary = {row[0]: (row[1], row[2]) for row in workbook["集計"].iter_rows(min_row=6, values_only=True)
               if row[0]}
    total, note = summary["会計用売上金額：判明分の合計（円）"]
    assert total == 77000 + 0 + 8000 + 13000
    assert "5 件中 4 件の合計" in note
    assert "1 件は値が不明" in note
    assert summary["会計用売上金額：値が不明（件）"][0] == 1
    assert "消費税の納税義務：免税事業者" in summary
    assert "経理方式：税込経理" in summary
    assert "インボイス登録：なし" in summary
