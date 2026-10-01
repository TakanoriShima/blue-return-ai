"""売上計上日の人間判断（data/decisions/revenue_dates.csv）の反映テスト。すべて一時フォルダ上の架空データ。"""

from __future__ import annotations

import csv
import json

import pytest
from openpyxl import load_workbook

from blue_return_ai import sales_workbook
from blue_return_ai.annual import AnnualError, load_revenue_date_decisions, run
from blue_return_ai.sample_template_c import write_sample as write_c

CONFIG = {
    "version": 1,
    "sources": {"invoices": "data/invoices"},
    "decisions_dir": "data/decisions",
    "processors": {"invoices": {"templates": ["template_c"]}},
}
HEADER = "source_key,revenue_date,revenue_date_basis,note"


@pytest.fixture
def repo(tmp_path):
    repo = tmp_path / "repo"
    invoices = repo / "data" / "invoices"
    write_c(invoices / "a.xlsx", pattern="A")
    write_c(invoices / "c.xlsx", pattern="C")
    config = repo / "data" / "config" / "annual.json"
    config.parent.mkdir(parents=True)
    config.write_text(json.dumps(CONFIG), encoding="utf-8")
    return repo


def run_repo(repo):
    return run(2026, repo / "data" / "config" / "annual.json", repo / "data" / "output" / "annual",
               repo_root=repo)


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def source_keys(repo):
    """判断ファイルを書くために、一度実行して source_key を得る（人間が sales.csv から写す操作に相当）。"""
    sales = read_csv(run_repo(repo)["run_dir"] / "sales.csv")
    return {row["template_pattern"]: row["source_key"] for row in sales}


def write_decisions(repo, *lines):
    path = repo / "data" / "decisions" / "revenue_dates.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join([HEADER, *lines]) + "\n", encoding="utf-8")
    return path


def sales_by_pattern(result):
    return {row["template_pattern"]: row for row in read_csv(result["run_dir"] / "sales.csv")}


def test_without_decisions_file_behavior_is_unchanged(repo):
    result = run_repo(repo)
    sales = sales_by_pattern(result)
    for row in sales.values():
        assert row["revenue_date"] == ""
        assert row["revenue_date_status"] == "unconfirmed"
        assert row["revenue_date_basis"] == ""
    assert result["manifest"]["processing"]["unresolved"]["by_code"]["REVENUE_DATE_UNCONFIRMED"] == 2
    assert result["manifest"]["processing"]["sales"]["revenue_date_confirmed"] == 0


def test_confirmed_decision_is_applied(repo):
    keys = source_keys(repo)
    write_decisions(repo, f"{keys['A']},2026-04-30,service_period_end,架空のメモ")
    result = run_repo(repo)
    sales = sales_by_pattern(result)
    assert sales["A"]["revenue_date"] == "2026-04-30"
    assert sales["A"]["revenue_date_status"] == "human_confirmed"
    assert sales["A"]["revenue_date_basis"] == "service_period_end"
    # 判断の無い売上は従来どおり
    assert sales["C"]["revenue_date"] == ""
    assert sales["C"]["revenue_date_status"] == "unconfirmed"
    # 原資料の日付は変更しない
    assert sales["A"]["document_date"] == "2026-05-31"

    items = read_csv(result["run_dir"] / "unresolved_items.csv")
    unconfirmed = {i["source_key"] for i in items if i["issue_code"] == "REVENUE_DATE_UNCONFIRMED"}
    assert keys["A"] not in unconfirmed
    assert keys["C"] in unconfirmed
    assert result["manifest"]["processing"]["sales"]["revenue_date_confirmed"] == 1
    # メモは出力に含めない
    assert "架空のメモ" not in (result["run_dir"] / "sales.csv").read_text(encoding="utf-8-sig")


def test_excel_shows_revenue_date_and_basis(repo):
    keys = source_keys(repo)
    write_decisions(repo, f"{keys['C']},2026-06-15,contract_based,")
    result = run_repo(repo)
    workbook = load_workbook(result["run_dir"] / sales_workbook.workbook_file_name(2026))
    sheet = workbook["売上一覧"]
    header = [c.value for c in sheet[3]]
    rows = [dict(zip(header, values)) for values in sheet.iter_rows(min_row=4, values_only=True)]
    confirmed = next(r for r in rows if r["帳票パターン"] == "C")
    assert confirmed["売上計上日"].date().isoformat() == "2026-06-15"
    assert confirmed["売上計上日の状態"] == "human_confirmed"
    assert confirmed["売上計上日の根拠"] == "contract_based"
    other = next(r for r in rows if r["帳票パターン"] == "A")
    assert other["売上計上日"] is None
    assert other["売上計上日の根拠"] is None
    summary = {row[0]: row[1] for row in workbook["集計"].iter_rows(min_row=6, values_only=True) if row[0]}
    assert summary["売上計上日 確定（件）"] == 1
    assert summary["売上計上日 未確定（件）"] == 1


def test_blank_revenue_date_is_not_confirmed(repo):
    keys = source_keys(repo)
    write_decisions(repo, f"{keys['A']},,service_period_end,未判断")
    sales = sales_by_pattern(run_repo(repo))
    assert sales["A"]["revenue_date_status"] == "unconfirmed"


def test_unused_decision_is_counted(repo):
    keys = source_keys(repo)
    write_decisions(repo, f"{keys['A']},2026-04-30,other,", "sales-1-" + "0" * 32 + ",2026-05-01,other,")
    result = run_repo(repo)
    assert result["manifest"]["processing"]["sales"]["revenue_date_confirmed"] == 1
    assert result["manifest"]["processing"]["sales"]["revenue_date_decisions_unused"] == 1


@pytest.mark.parametrize(
    ("line", "message"),
    [
        ("{key},2026/04/30,service_period_end,", "YYYY-MM-DD"),          # 日付形式不正
        ("{key},2026-02-30,service_period_end,", "存在しない日付"),
        ("{key},2025-12-31,service_period_end,", "対象年と異なります"),   # 対象年不一致
        ("{key},2026-04-30,invoice_date,", "revenue_date_basis"),        # 未知の根拠
        ("{key},2026-04-30,,", "revenue_date_basis"),                    # 根拠なし
        (",2026-04-30,other,", "source_key が空"),
    ],
)
def test_invalid_decisions_are_rejected(repo, line, message):
    keys = source_keys(repo)
    write_decisions(repo, line.format(key=keys["A"]))
    with pytest.raises(AnnualError) as error:
        run_repo(repo)
    assert message in str(error.value)
    assert "2 行目" in str(error.value)


def test_duplicate_source_key_is_rejected(repo):
    keys = source_keys(repo)
    write_decisions(repo, f"{keys['A']},2026-04-30,other,", f"{keys['A']},2026-05-31,other,")
    with pytest.raises(AnnualError) as error:
        run_repo(repo)
    assert "重複" in str(error.value)
    assert "3 行目" in str(error.value)


def test_errors_stop_before_creating_output(repo):
    keys = source_keys(repo)
    out_root = repo / "data" / "output" / "annual"
    before = set(out_root.iterdir())
    write_decisions(repo, f"{keys['A']},2025-01-01,other,")
    with pytest.raises(AnnualError):
        run_repo(repo)
    assert set(out_root.iterdir()) == before


def test_missing_columns_are_rejected(tmp_path):
    path = tmp_path / "revenue_dates.csv"
    path.write_text("source_key,revenue_date\nx,2026-01-01\n", encoding="utf-8")
    with pytest.raises(AnnualError):
        load_revenue_date_decisions(tmp_path, 2026)


def test_loader_without_file_or_dir_returns_empty(tmp_path):
    assert load_revenue_date_decisions(None, 2026) == {}
    assert load_revenue_date_decisions(tmp_path, 2026) == {}


def test_error_messages_do_not_contain_values(repo):
    keys = source_keys(repo)
    write_decisions(repo, f"{keys['A']},2025-12-31,service_period_end,架空のメモ")
    with pytest.raises(AnnualError) as error:
        run_repo(repo)
    text = str(error.value)
    for value in (keys["A"], "2025-12-31", "架空のメモ"):
        assert value not in text
