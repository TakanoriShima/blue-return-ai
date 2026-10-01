"""Excel ファイルの読み取り（docs/data_model.md §11）。

- 元ファイルはバイト列として一度だけ読み込み、以降はメモリ上で扱う。
  元ファイルを開いたまま操作したり、保存したりしない。
- openpyxl は数式を計算しないため、数式セルかどうか（data_only=False）と
  キャッシュ値（data_only=True）の両方を取得する。
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class SheetNotFoundError(Exception):
    """テンプレートが指定するシートが存在しない。"""


@dataclass(frozen=True)
class CellValue:
    value: Any  # data_only=True で得た値。数式セルの場合は最後に保存されたキャッシュ値
    is_formula: bool

    @property
    def missing_cached_value(self) -> bool:
        return self.is_formula and self.value is None


EMPTY_CELL = CellValue(value=None, is_formula=False)


@dataclass(frozen=True)
class ExcelSource:
    source_hash: str
    cells: dict[str, CellValue]

    def cell(self, coordinate: str) -> CellValue:
        return self.cells.get(coordinate, EMPTY_CELL)


def sha256_of_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def read_excel(path: Path, sheet_name: str) -> ExcelSource:
    data = Path(path).read_bytes()
    formula_flags = _read_cells(data, sheet_name, data_only=False)
    values = _read_cells(data, sheet_name, data_only=True)
    cells = {
        coordinate: CellValue(value=values.get(coordinate), is_formula=is_formula)
        for coordinate, is_formula in formula_flags.items()
    }
    return ExcelSource(source_hash=sha256_of_bytes(data), cells=cells)


def _read_cells(data: bytes, sheet_name: str, *, data_only: bool) -> dict[str, Any]:
    """data_only=False のときは「数式セルか」を、True のときは値を返す。"""
    workbook = load_workbook(BytesIO(data), data_only=data_only)
    try:
        if sheet_name not in workbook.sheetnames:
            # シート名は公開テンプレート定義の値であり、実データではない
            raise SheetNotFoundError(f"シート '{sheet_name}' が見つかりません")
        sheet = workbook[sheet_name]
        result: dict[str, Any] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if data_only:
                    result[cell.coordinate] = cell.value
                else:
                    result[cell.coordinate] = cell.data_type == "f"
        return result
    finally:
        workbook.close()
