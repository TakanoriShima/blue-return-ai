"""template_instructor 形式（シート「講師」）の「完全な架空」Excel 請求書を生成する（開発・テスト用）。

- 氏名・会社名・銀行名・口座番号・商品名・金額はすべて架空の値であり、実在のものを模倣していない。
- ラベルと列の構成は template_instructor の定義から作る（実際の帳票の値は使わない）。
- 金額は数式ではなく値として書き込む。既存ファイルは上書きしない。data/ 配下へは書き込まない。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from . import template_instructor as t

REPO_ROOT = Path(__file__).resolve().parents[2]
HEADER_ROW = 20

# 架空の明細: (商品名, 数量, 単位, 単価, 税率, 金額)
SAMPLE_ITEMS: list[tuple[str, Any, str, int, Any, int]] = [
    ("架空講座A", 2, "時間", 5000, "10%", 10000),
    ("架空講座B", 1, "回", 3000, "10%", 3000),
]
# 架空の集計欄（小計 13000、消費税 1300、源泉税額 1327、合計 12973）
SAMPLE_SUMMARY: dict[str, Any] = {
    "net_amount": 13000,
    "tax_amount": 1300,
    "withholding_tax": 1327,
    "document_total": 12973,
}


def build_workbook(
    items: list[tuple] | None = None,
    summary: dict[str, Any] | None = None,
    overrides: dict[str, Any] | None = None,
    blank_rows_after_items: int = 0,
    include_breakdown: bool = True,
    document_date: Any = datetime(2026, 6, 30),
) -> Workbook:
    items = SAMPLE_ITEMS if items is None else items
    summary = {**SAMPLE_SUMMARY, **(summary or {})}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = t.SHEET_NAME
    sheet["A1"] = "御請求書"
    sheet["A3"] = "架空 研修事業部 御中"
    sheet["G3"] = t.DOCUMENT_DATE_LABEL
    sheet["H3"] = document_date
    if isinstance(document_date, datetime):
        sheet["H3"].number_format = "yyyy/mm/dd"
    sheet["A8"] = "架空 講師"
    sheet["A10"] = "振込先：架空銀行 架空支店 普通 0000000"

    for key, (label, column) in t.LINE_HEADERS.items():
        sheet[f"{column}{HEADER_ROW}"] = label
    sheet.merge_cells(f"A{HEADER_ROW}:C{HEADER_ROW}")

    row = HEADER_ROW + 1
    for name, quantity, unit, unit_price, tax_rate, amount in items:
        sheet[f"A{row}"] = name
        sheet.merge_cells(f"A{row}:C{row}")
        sheet[f"D{row}"] = quantity
        sheet[f"E{row}"] = unit
        sheet[f"F{row}"] = unit_price
        sheet[f"G{row}"] = tax_rate
        if isinstance(tax_rate, float):
            sheet[f"G{row}"].number_format = "0%"
        sheet[f"H{row}"] = amount
        row += 1
    row += blank_rows_after_items

    for key, label in t.SUMMARY_LABELS.items():
        sheet[f"{t.SUMMARY_LABEL_COLUMN}{row}"] = label
        sheet.merge_cells(f"G{row}:H{row}")
        if summary.get(key) is not None:
            sheet[f"G{row}"] = summary[key]
        row += 1

    if include_breakdown:
        row += 1
        sheet[f"A{row}"] = "税率別内訳"
        sheet[f"D{row}"] = "税抜金額"
        sheet[f"F{row}"] = "消費税額"
        sheet[f"A{row + 1}"] = "10%対象"
        sheet[f"D{row + 1}"] = summary["net_amount"]
        sheet[f"F{row + 1}"] = summary["tax_amount"]
        row += 2
    sheet[f"A{row + 1}"] = "※ この請求書は blue-return-ai の開発・テスト用に作成した完全な架空データです。"

    for coordinate, value in (overrides or {}).items():
        sheet[coordinate] = value
    workbook.properties.creator = "blue-return-ai sample generator"
    return workbook


def write_sample(path: Path, **kwargs) -> Path:
    path = Path(path)
    if path.resolve().is_relative_to(REPO_ROOT / "data"):
        raise ValueError("架空データを data/ 配下へ書き込むことはできません")
    if path.exists():
        raise FileExistsError("既存ファイルは上書きしません")
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(**kwargs).save(path)
    return path
