"""template_c 形式（パターン A〜E）の「完全な架空」Excel 請求書を生成する（開発・テスト用）。

    python -m blue_return_ai.sample_template_c sample_data/invoices/template_c_sample.xlsx

- 取引先名・品目・金額・日付・氏名はすべて架空の値であり、実在の企業・人物・請求書を模倣していない。
- ラベル構成は template_c.PATTERNS の定義から作る（実際の帳票の値は使わない）。
- 明細の件数によって集計欄の行位置が変わる書式を再現する（items 引数で件数を変えられる）。
- 金額は数式ではなく値として書き込む（openpyxl で作成した数式セルはキャッシュ値を持たないため）。
- 既存ファイルは上書きしない。data/ 配下へは書き込まない。
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from . import template_c

REPO_ROOT = Path(__file__).resolve().parents[2]

# パターンごとの架空データ: (明細 [(品目, 数量, 単価, 金額)], 集計欄 {項目: 値}, ご請求金額)
SAMPLES: dict[str, tuple[list[tuple[str, Any, int, int]], dict[str, int], int]] = {
    "A": ([("架空研修講師業務", 3, 20000, 60000), ("架空教材作成業務", 1, 10000, 10000)],
          {"net_amount": 70000, "tax_amount": 7000, "gross_amount": 77000,
           "withholding_tax": 7147, "expected_payment_amount": 69853}, 77000),
    "B": ([("架空講師業務", 10, 3300, 33000)],
          {"gross_amount": 33000, "withholding_base": 30000, "withholding_tax": 3063,
           "expected_payment_amount": 29937}, 33000),
    "C": ([("架空相談業務", 2, 5000, 10000)],
          {"net_amount": 10000, "tax_amount": 1000, "gross_amount": 11000}, 11000),
    "D": ([("架空作業", 1, 8000, 8000)], {"document_subtotal": 8000}, 8000),
    "E": ([("架空講義", 4, 2750, 11000)],
          {"gross_amount": 11000, "net_amount": 10000, "tax_amount": 1000,
           "withholding_tax": 1021, "expected_payment_amount": 9979}, 11000),
}
HEADER_ROW = 13
HEADER_COLUMNS = {"description": "B", "quantity": "J", "unit_price": "L", "amount": "N"}
SUMMARY_LABEL_COLUMN = "J"
SUMMARY_VALUE_COLUMN = "L"


def _pattern(name: str) -> template_c.Pattern:
    return next(p for p in template_c.PATTERNS if p.name == name)


def build_workbook(
    pattern: str = "A",
    items: list[tuple[str, Any, int, int]] | None = None,
    summary: dict[str, Any] | None = None,
    overrides: dict[str, Any] | None = None,
    blank_rows_after_items: int = 0,
    document_date: Any = datetime(2026, 5, 31),
) -> Workbook:
    """架空の template_c 請求書を作る。

    items: (品目, 数量, 単価, 金額) の一覧。summary: 集計欄の値（None で空欄）。
    document_date: 請求日セルの値（datetime・シリアル値・文字列）。
    overrides: 最後にセルの値を差し替える（None で空欄にする）。
    """
    definition = _pattern(pattern)
    default_items, default_summary, document_total = SAMPLES[pattern]
    items = default_items if items is None else items
    summary = {**default_summary, **(summary or {})}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = template_c.SHEET_NAME
    sheet["B1"] = "請　求　書"
    sheet["B3"] = "架空サンプル研修株式会社"
    sheet["K3"] = template_c.DOCUMENT_DATE_LABEL
    sheet["L3"] = document_date
    if isinstance(document_date, datetime):
        sheet["L3"].number_format = "yyyy/mm/dd"
    sheet["B10"] = template_c.DOCUMENT_TOTAL_LABEL
    sheet["E10"] = document_total
    sheet["I10"] = "（税込）"
    sheet["L10"] = "名前："
    sheet["M10"] = "架空 請求者"

    for key, column in HEADER_COLUMNS.items():
        sheet[f"{column}{HEADER_ROW}"] = definition.line_headers[key]
    sheet.merge_cells(f"B{HEADER_ROW}:I{HEADER_ROW}")

    row = HEADER_ROW + 1
    for description, quantity, unit_price, amount in items:
        sheet[f"B{row}"] = description
        sheet[f"J{row}"] = quantity
        sheet[f"L{row}"] = unit_price
        sheet[f"N{row}"] = amount
        sheet.merge_cells(f"B{row}:I{row}")
        row += 1
    row += blank_rows_after_items

    for key, label in definition.summary.items():
        sheet[f"{SUMMARY_LABEL_COLUMN}{row}"] = label
        if summary.get(key) is not None:
            sheet[f"{SUMMARY_VALUE_COLUMN}{row}"] = summary[key]
            sheet[f"{SUMMARY_VALUE_COLUMN}{row}"].number_format = "#,##0"
        row += 1
    sheet[f"B{row + 1}"] = "※ この請求書は blue-return-ai の開発・テスト用に作成した完全な架空データです。"

    for coordinate, value in (overrides or {}).items():
        sheet[coordinate] = value
    workbook.properties.creator = "blue-return-ai sample generator"
    return workbook


def write_sample(path: Path, **kwargs) -> Path:
    path = Path(path)
    if path.resolve().is_relative_to(REPO_ROOT / "data"):
        raise ValueError("架空データを data/ 配下へ書き込むことはできません")
    if path.exists():
        raise FileExistsError("既存ファイルは上書きしません")
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(**kwargs).save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if len(args) != 1:
        print("使い方: python -m blue_return_ai.sample_template_c <出力先.xlsx>", file=sys.stderr)
        return 2
    write_sample(Path(args[0]))
    print(f"架空の template_c 請求書を作成しました: {args[0]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
