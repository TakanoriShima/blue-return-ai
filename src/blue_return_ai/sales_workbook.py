"""年間処理の売上データから、人間確認用の Excel「<年>年 売上・請求書集計（暫定）」を作る（docs/data_model.md §17.6）。

- 確定した売上帳・青色申告決算書ではない。売上計上日などの未確定事項がある前提の、確認用の資料。
- 値は年間処理の出力（sales.csv / sales_line_items.csv / unresolved_items.csv と同じデータ）だけを使い、
  データモデルに無い値を推測して追加しない。売上計上日が未確定なら空欄のまま（請求日で補完しない）。
- 金額の合計は「判明している値だけ」の合計とし、不明（null）を 0 として扱わない。不明の件数を必ず併記する。
- 文字列は常に文字列として書き込み、Excel で数式として解釈されないようにする。
- 出力先は data/ 配下の実行フォルダ。既存ファイルは上書きしない。
"""

from __future__ import annotations

from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AMOUNT_FORMAT = "#,##0"
DATE_FORMAT = "yyyy-mm-dd"
HEADER_ROW = 3
NOTICE = "※ 暫定資料です。確定した売上帳・青色申告決算書ではありません。売上計上日などに未確定の事項があります。"

# 売上一覧：(見出し, sales.csv の列, 種類)
SALES_SHEET_COLUMNS: list[tuple[str, str, str]] = [
    ("資料日付（請求日等）", "document_date", "date"),
    ("売上計上日", "revenue_date", "date"),
    ("売上計上日の状態", "revenue_date_status", "text"),
    ("業務期間（開始）", "service_period_from", "date"),
    ("業務期間（終了）", "service_period_to", "date"),
    ("業務対象月", "service_period_month", "text"),
    ("取引先", "customer", "text"),
    ("内容", "description", "text"),
    ("税抜金額", "net_amount", "amount"),
    ("消費税額", "tax_amount", "amount"),
    ("税込金額", "gross_amount", "amount"),
    ("源泉徴収税額", "withholding_tax", "amount"),
    ("その他控除", "deductions_total", "amount"),
    ("振込予定額", "expected_payment_amount", "amount"),
    ("会計用売上金額（税込経理）", "accounting_sales_amount", "amount"),
    ("会計用売上金額の根拠", "accounting_sales_amount_basis", "text"),
    ("支払期限", "payment_due", "date"),
    ("税区分", "tax_treatment", "text"),
    ("税率(%)", "tax_rate", "number"),
    ("源泉徴収の状態", "withholding_status", "text"),
    ("計算で求めた項目", "calculated_fields", "text"),
    ("確認状態（review_status）", "review_status", "text"),
    ("警告", "warnings", "text"),
    ("テンプレート", "source_name", "text"),
    ("帳票パターン", "template_pattern", "text"),
    ("ご請求金額（帳票の値・確認用）", "document_total", "amount"),
    ("小計（意味未確定・確認用）", "document_subtotal", "amount"),
    ("源泉徴収の対象額（確認用）", "withholding_base", "amount"),
    ("明細件数", "line_item_count", "number"),
    ("record_id", "record_id", "text"),
    ("source_key", "source_key", "text"),
    ("document_key", "document_key", "text"),
]

LINE_SHEET_COLUMNS: list[tuple[str, str, str]] = [
    ("売上No", "sales_no", "number"),
    ("明細No", "line_no", "number"),
    ("資料日付", "document_date", "date"),
    ("取引先", "customer", "text"),
    ("品目", "description", "text"),
    ("数量", "quantity", "decimal"),
    ("単位", "unit", "text"),
    ("単価", "unit_price", "amount"),
    ("金額", "amount", "amount"),
    ("税率(%)", "tax_rate", "number"),
    ("record_id", "record_id", "text"),
]

ISSUE_SHEET_COLUMNS: list[tuple[str, str, str]] = [
    ("要確認ID", "item_id", "text"),
    ("売上No", "sales_no", "number"),
    ("警告コード", "issue_code", "text"),
    ("重大度", "severity", "text"),
    ("対象項目", "field", "text"),
    ("内容", "message", "text"),
    ("判断（記入欄）", "decision_status", "text"),
    ("判断した値（記入欄）", "decision_value", "text"),
    ("メモ（記入欄）", "decision_note", "text"),
    ("source_key", "source_key", "text"),
    ("document_key", "document_key", "text"),
]

# 集計する金額項目：(見出し, sales.csv の列)
SUMMARY_AMOUNTS: list[tuple[str, str]] = [
    ("税抜金額", "net_amount"),
    ("消費税額", "tax_amount"),
    ("税込金額", "gross_amount"),
    ("源泉徴収税額", "withholding_tax"),
    ("その他控除", "deductions_total"),
    ("振込予定額", "expected_payment_amount"),
    ("会計用売上金額", "accounting_sales_amount"),
]

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_TITLE_FONT = Font(bold=True, size=14)


def workbook_file_name(year: int) -> str:
    return f"sales_summary_{year}_provisional.xlsx"


def _set(cell, value: Any, kind: str) -> None:
    """値を種類に応じて書き込む。None は空欄（0 にしない）。文字列は数式として解釈させない。"""
    if value is None or value == "":
        return
    if kind == "amount":
        cell.value = value
        cell.number_format = AMOUNT_FORMAT
    elif kind == "number":
        cell.value = value
    elif kind == "decimal":
        try:
            cell.value = Decimal(str(value))
        except InvalidOperation:
            _set_text(cell, str(value))
    elif kind == "date":
        try:
            cell.value = date.fromisoformat(str(value))
            cell.number_format = DATE_FORMAT
        except ValueError:
            _set_text(cell, str(value))
    else:
        _set_text(cell, value if isinstance(value, str) else str(value))


def _set_text(cell, text: str) -> None:
    cell.value = text
    cell.data_type = "s"  # 「=」で始まる文字列も数式にしない


def _write_table(sheet, title: str, columns: list[tuple[str, str, str]], rows: list[dict]) -> None:
    sheet["A1"] = title
    sheet["A1"].font = _TITLE_FONT
    _set_text(sheet["A2"], NOTICE)
    for index, (header, _, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index)
        _set_text(cell, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[get_column_letter(index)].width = max(10, min(28, len(header) * 2 + 2))
    for row_number, row in enumerate(rows, start=HEADER_ROW + 1):
        for index, (_, key, kind) in enumerate(columns, start=1):
            _set(sheet.cell(row=row_number, column=index), row.get(key), kind)
    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=1)
    if rows:
        last = get_column_letter(len(columns))
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{last}{HEADER_ROW + len(rows)}"


ACCOUNTING_LABELS = {
    "tax_status": ("消費税の納税義務", {"exempt": "免税事業者", "taxable": "課税事業者"}),
    "invoice_registration": ("インボイス登録", {True: "あり", False: "なし"}),
    "consumption_tax_accounting": ("経理方式", {"inclusive": "税込経理", "exclusive": "税抜経理"}),
}


def build_summary_rows(sales: list[dict], invoice_documents: Counter,
                       accounting: dict | None = None) -> list[list[Any]]:
    """集計シートの行（[項目, 値, 補足]）。null を 0 として扱わない。"""
    total = len(sales)
    rows: list[list[Any]] = [["■ 会計上の前提（設定ファイルの年度別設定）", None, None]]
    if accounting:
        for key, (label, names) in ACCOUNTING_LABELS.items():
            rows.append([f"{label}：{names.get(accounting.get(key), accounting.get(key))}", None, None])
    else:
        rows.append(["この年の会計設定がありません", None, "会計用売上金額は決められないため、すべて不明になります"])
    rows += [
        ["■ 取込状況", None, None],
        ["取り込み済み売上資料（件）", total, "Excel 請求書から取り込めた資料。PDF からは売上を作っていない"],
    ]
    for status in ("imported", "manual", "cross_check_only", "unresolved", "unsupported"):
        rows.append([f"請求書フォルダの資料：{status}（件）", invoice_documents.get(status, 0), None])

    rows.append(["■ 確認状態", None, None])
    review = Counter(s.get("review_status") for s in sales)
    for status, note in (("unreviewed", "機械検証で問題なし・人間は未確認"),
                         ("needs_review", "警告・欠落等があり人間の確認が必要"),
                         ("reviewed", "人間が確認済み（プログラムは設定しない）")):
        rows.append([f"review_status = {status}（件）", review.get(status, 0), note])
    confirmed = sum(1 for s in sales if s.get("revenue_date"))
    rows.append(["売上計上日 確定（件）", confirmed, None])
    rows.append(["売上計上日 未確定（件）", total - confirmed, "請求日を売上計上日として自動で補完していない"])

    rows.append(["■ 金額（判明している値だけの合計。不明な値を 0 として扱っていない）", None, None])
    for label, key in SUMMARY_AMOUNTS:
        known = [s[key] for s in sales if isinstance(s.get(key), int) and not isinstance(s.get(key), bool)]
        unknown = total - len(known)
        if not total:
            note = "対象なし"
        elif unknown:
            note = f"{total} 件中 {len(known)} 件の合計。{unknown} 件は値が不明のため含まない（不完全な合計）"
        else:
            note = f"{total} 件すべての値が判明"
        rows.append([f"{label}：判明分の合計（円）", sum(known) if known else None, note])
        rows.append([f"{label}：値が不明（件）", unknown, None])

    rows.append(["■ 内訳（件数）", None, None])
    for label, key in (("帳票パターン", "template_pattern"), ("テンプレート", "source_name"),
                       ("税区分", "tax_treatment"), ("源泉徴収の状態", "withholding_status")):
        for value, count in sorted(Counter(s.get(key) or "（なし）" for s in sales).items()):
            rows.append([f"{label} = {value}（件）", count, None])
    return rows


def build_workbook(sales: list[dict], line_items: list[dict], issues: list[dict],
                   invoice_documents: Counter, year: int, run_id: str, created_at: str,
                   accounting: dict | None = None) -> Workbook:
    """sales：sales.csv の行（dict）、line_items：sales_line_items.csv の行、issues：売上の要確認。"""
    title = f"{year}年 売上・請求書集計（暫定）"
    sales_no = {row["record_id"]: number for number, row in enumerate(sales, start=1)}
    by_key = {row["source_key"]: number for number, row in enumerate(sales, start=1)}
    record = {row["record_id"]: row for row in sales}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "集計"
    summary["A1"] = title
    summary["A1"].font = _TITLE_FONT
    _set_text(summary["A2"], NOTICE)
    _set_text(summary["A3"], f"実行ID：{run_id} ／ 作成日時：{created_at}")
    for index, header in enumerate(("項目", "値", "補足"), start=1):
        cell = summary.cell(row=5, column=index)
        _set_text(cell, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for row_number, (label, value, note) in enumerate(build_summary_rows(sales, invoice_documents, accounting), start=6):
        _set_text(summary.cell(row=row_number, column=1), label)
        if label.startswith("■"):
            summary.cell(row=row_number, column=1).font = _HEADER_FONT
        if value is not None:
            cell = summary.cell(row=row_number, column=2)
            cell.value = value
            cell.number_format = AMOUNT_FORMAT
        if note:
            _set_text(summary.cell(row=row_number, column=3), note)
    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 70

    numbered = [{**row, "sales_no": number} for number, row in enumerate(sales, start=1)]
    _write_table(workbook.create_sheet("売上一覧"), f"{title}：売上一覧（1 行 = 1 売上資料）",
                 [("売上No", "sales_no", "number"), *SALES_SHEET_COLUMNS], numbered)

    lines = [
        {**line, "sales_no": sales_no.get(line["record_id"]),
         "document_date": record.get(line["record_id"], {}).get("document_date"),
         "customer": record.get(line["record_id"], {}).get("customer")}
        for line in line_items
    ]
    _write_table(workbook.create_sheet("売上明細"), f"{title}：売上明細（売上No で売上一覧と対応）",
                 LINE_SHEET_COLUMNS, lines)

    issue_rows = [{**issue, "sales_no": by_key.get(issue.get("source_key"))} for issue in issues]
    _write_table(workbook.create_sheet("要確認"),
                 f"{title}：要確認（売上に関するもの。記入欄は空欄で出力）", ISSUE_SHEET_COLUMNS, issue_rows)

    workbook.properties.creator = "blue-return-ai"
    workbook.properties.title = title
    return workbook


def write_workbook(path: Path, workbook: Workbook) -> Path:
    """既存ファイルは上書きしない。"""
    buffer = BytesIO()
    workbook.save(buffer)
    with Path(path).open("xb") as f:
        f.write(buffer.getvalue())
    return Path(path)
